#!/usr/bin/env python3
"""Generate and evaluate a simplified manual audit for multilingual RDF-to-text data.

The audit uses one high-level human judgement per sub-QA:

* QA1.1: structural/linkage risk;
* QA1.2: recurring-component inconsistency risk;
* QA2.1: source-knowledge preservation risk;
* QA2.2: target RDF--text faithfulness risk;
* QA3.1: target-language/leakage risk; and
* QA3.2: risk that divergence is not legitimate variation.

Human labels are deliberately simple:

    Low | Medium | High | Cannot judge

The annotator workbook is blinded: it does not show the automatic risk level or
metric values.  A separate key workbook stores the automatic risk and its basis.
The ``evaluate`` command compares the completed human labels with the automatic
risk strata and produces agreement and review-priority statistics.

Examples
--------
Generate the audit files::

    python manual_audit_pipeline_reproducible.py generate \
        --results-root new_results \
        --output-dir manual_audit

Evaluate a completed workbook::

    python manual_audit_pipeline_reproducible.py evaluate \
        --annotated manual_audit/manual_audit_annotator_completed.xlsx \
        --key-csv manual_audit/manual_audit_key.csv \
        --output-dir manual_audit/evaluation
"""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import math
import random
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:
    raise ImportError(
        "This script requires openpyxl. Install it with: "
        "python -m pip install openpyxl"
    ) from exc

RISK_LEVELS = ("Low", "Medium", "High")
MANUAL_LEVELS = ("Low", "Medium", "High", "Cannot judge")
RISK_TO_NUMBER = {"Low": 0, "Medium": 1, "High": 2}

SUBQA_ORDER = ("QA1.1", "QA1.2", "QA2.1", "QA2.2", "QA3.1", "QA3.2")
SHEET_NAMES = {
    "QA1.1": "QA1_1",
    "QA1.2": "QA1_2",
    "QA2.1": "QA2_1",
    "QA2.2": "QA2_2",
    "QA3.1": "QA3_1",
    "QA3.2": "QA3_2",
}

SUBQA_DEFINITIONS = {
    "QA1.1": {
        "title": "Structural and instance integrity",
        "question": (
            "Is there a structural or linkage problem that makes this "
            "RDF-to-text entry unreliable?"
        ),
        "low": "All required RDF, identifier, and text components appear complete and correctly linked.",
        "medium": "A possible structural or linkage issue is present, but its impact is limited or uncertain.",
        "high": "A clear missing, duplicated, malformed, or incorrectly linked component makes the entry unreliable.",
        "example": "A missing target text, duplicated lexicalisation ID, malformed triple, or mismatched source/target ID.",
    },
    "QA1.2": {
        "title": "Recurring-component consistency",
        "question": (
            "Does this recurring RDF component show a harmful inconsistency "
            "rather than acceptable target-language variation?"
        ),
        "low": "The mapping is stable, or the observed alternatives are clearly legitimate.",
        "medium": "The variation is review-worthy but may reflect an alias, inflection, or conventional alternative.",
        "high": "The component is represented by clearly incompatible target forms that can change or obscure its meaning.",
        "example": "A source entity repeatedly maps to two incompatible entities, rather than to an alias or scientific/common name pair.",
    },
    "QA2.1": {
        "title": "Source-to-target knowledge preservation",
        "question": (
            "Does the Spanish or Catalan adaptation fail to preserve knowledge "
            "contained in the aligned English source?"
        ),
        "low": "The target triples and text preserve the source entities, relations, and factual values.",
        "medium": "A possible semantic divergence exists, but it may be explained by translation or lexical variation.",
        "high": "A clear entity, relation, literal, or record-level meaning change is present.",
        "example": "A date, number, entity, or relation differs from the aligned English source.",
    },
    "QA2.2": {
        "title": "Target RDF--text faithfulness",
        "question": (
            "Does the target verbalisation omit information from its target RDF "
            "or introduce unsupported information?"
        ),
        "low": "The target text expresses the RDF facts without a clear unsupported addition.",
        "medium": "A fact may be implicit, weakly expressed, or potentially unsupported, but the case is ambiguous.",
        "high": "A clear RDF fact is missing, a factual value is altered, or unsupported content is added.",
        "example": "One RDF relation is absent from the text, or the text adds a fact not licensed by the target triples.",
    },
    "QA3.1": {
        "title": "Target-language validity and English leakage",
        "question": (
            "Is the target lexicalisation in the wrong language or does it "
            "contain substantial English leakage?"
        ),
        "low": "The text is acceptable Spanish or Catalan and contains no substantial English leakage.",
        "medium": "The language is uncertain or contains a short fragment that may be a name, technical expression, or leakage.",
        "high": "The text is clearly in the wrong language, copied from English, or contains substantial untranslated English material.",
        "example": "A complete English copy or a clearly untranslated English clause inside an otherwise target-language sentence.",
    },
    "QA3.2": {
        "title": "Legitimate linguistic and reference variation",
        "question": (
            "Is the divergence from the English reference unlikely to be "
            "legitimate linguistic or multi-reference variation?"
        ),
        "low": "The difference is a plausible translation, paraphrase, aggregation, or reference-selection effect.",
        "medium": "The divergence is unusual and warrants review, but may still be a valid alternative.",
        "high": "The divergence is difficult to justify as legitimate variation and appears to alter or distort the intended content.",
        "example": "The target text is much closer to an unrelated reference or is abnormally short/long in a way that loses content.",
    },
}

ANNOTATOR_HEADERS = [
    "Case ID",
    "Language",
    "Category",
    "Record or component",
    "English triples",
    "Target triples",
    "Aligned English text",
    "Target text",
    "Comparison or context",
    "Manual risk",
    "Manual notes",
]

KEY_EXTRA_HEADERS = [
    "Automatic risk",
    "Automatic basis",
    "Automatic signals",
]


@dataclass(frozen=True)
class AuditPaths:
    qa11: Path
    qa12_profiles: Path
    qa12_mappings: Path
    qa2: Path
    qa3: Path


@dataclass
class Case:
    case_id: str
    subqa: str
    language: str
    language_label: str
    category: str
    focus: str
    english_triples: str
    target_triples: str
    english_text: str
    target_text: str
    comparison_context: str
    automatic_risk: str
    automatic_basis: str
    automatic_signals: dict[str, Any]
    uniqueness_token: str

    def annotator_row(self) -> list[Any]:
        return [
            self.case_id,
            self.language_label,
            self.category,
            self.focus,
            self.english_triples,
            self.target_triples,
            self.english_text,
            self.target_text,
            self.comparison_context,
            "",
            "",
        ]

    def key_row(self) -> list[Any]:
        return self.annotator_row()[:-2] + [
            self.automatic_risk,
            self.automatic_basis,
            json.dumps(self.automatic_signals, ensure_ascii=False, sort_keys=True),
        ]

    def combined_key_dict(self) -> dict[str, Any]:
        return {
            "case_id": self.case_id,
            "subqa": self.subqa,
            "language": self.language,
            "language_label": self.language_label,
            "category": self.category,
            "record_or_component": self.focus,
            "english_triples": self.english_triples,
            "target_triples": self.target_triples,
            "aligned_english_text": self.english_text,
            "target_text": self.target_text,
            "comparison_or_context": self.comparison_context,
            "automatic_risk": self.automatic_risk,
            "automatic_basis": self.automatic_basis,
            "automatic_signals": json.dumps(
                self.automatic_signals,
                ensure_ascii=False,
                sort_keys=True,
            ),
        }


def normalize_space(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.casefold() in {"nan", "none", "null"}:
        return ""
    return re.sub(r"\s+", " ", text).strip()


def parse_bool(value: Any) -> bool | None:
    text = normalize_space(value).casefold()
    if text in {"true", "1", "yes"}:
        return True
    if text in {"false", "0", "no"}:
        return False
    return None


def parse_float(value: Any) -> float | None:
    text = normalize_space(value)
    if not text:
        return None
    try:
        number = float(text)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def parse_int(value: Any) -> int | None:
    number = parse_float(value)
    if number is None:
        return None
    return int(number)


def parse_list(value: Any) -> list[Any]:
    if isinstance(value, list):
        return value
    text = normalize_space(value)
    if not text:
        return []
    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(text)
            if isinstance(parsed, (list, tuple)):
                return list(parsed)
        except Exception:
            pass
    return [text]


def format_list(value: Any, *, limit: int | None = None) -> str:
    values = [normalize_space(item) for item in parse_list(value)]
    values = [item for item in values if item]
    if limit is not None and len(values) > limit:
        omitted = len(values) - limit
        values = values[:limit] + [f"… ({omitted} more)"]
    return "\n".join(values)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Required input file not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv_rows(path: Path, rows: Sequence[dict[str, Any]], fieldnames: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({name: row.get(name, "") for name in fieldnames})


def quantile(values: Iterable[float | None], q: float) -> float:
    usable = sorted(value for value in values if value is not None and math.isfinite(value))
    if not usable:
        return math.nan
    if len(usable) == 1:
        return usable[0]
    position = (len(usable) - 1) * q
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return usable[lower]
    weight = position - lower
    return usable[lower] * (1.0 - weight) + usable[upper] * weight


def stable_seed(seed: int, *parts: str) -> int:
    digest = hashlib.sha256("||".join(parts).encode("utf-8")).hexdigest()
    return seed + int(digest[:8], 16)


def discover_paths(results_root: Path) -> AuditPaths:
    candidates = {
        "qa11": [
            results_root / "QA1_stress_aligned" / "qa11_original_entry_metrics.csv",
            results_root / "qa11_original_entry_metrics.csv",
        ],
        "qa12_profiles": [
            results_root / "QA1_stress_aligned" / "qa12_mapping_profiles.csv",
            results_root / "qa12_mapping_profiles.csv",
        ],
        "qa12_mappings": [
            results_root / "QA1_stress_aligned" / "qa12_component_mappings.csv",
            results_root / "qa12_component_mappings.csv",
        ],
        "qa2": [
            results_root / "QA2_stress_aligned" / "qa2_lexicalisation_metrics.csv",
            results_root / "qa2_lexicalisation_metrics.csv",
        ],
        "qa3": [
            results_root / "QA3_stress_aligned" / "qa3_lexicalisation_metrics.csv",
            results_root / "qa3_lexicalisation_metrics.csv",
        ],
    }

    resolved: dict[str, Path] = {}
    for name, options in candidates.items():
        existing = next((path for path in options if path.exists()), None)
        if existing is None:
            raise FileNotFoundError(
                f"Could not locate {name}. Checked: "
                + ", ".join(str(path) for path in options)
            )
        resolved[name] = existing
    return AuditPaths(**resolved)


def first_lexicalisation_index(qa2_rows: Sequence[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    index: dict[tuple[str, str], dict[str, str]] = {}
    for row in qa2_rows:
        key = (normalize_space(row.get("record_key")), normalize_space(row.get("language")))
        index.setdefault(key, row)
    return index


def thresholds_by_language(
    rows: Sequence[dict[str, str]],
    columns: Sequence[str],
    quantiles: Sequence[float],
) -> dict[str, dict[str, dict[float, float]]]:
    grouped: dict[str, dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for row in rows:
        language = normalize_space(row.get("language"))
        for column in columns:
            value = parse_float(row.get(column))
            if value is not None:
                grouped[language][column].append(value)

    output: dict[str, dict[str, dict[float, float]]] = defaultdict(dict)
    for language, column_values in grouped.items():
        for column, values in column_values.items():
            output[language][column] = {
                q: quantile(values, q)
                for q in quantiles
            }
    return dict(output)


def make_case_id(subqa: str, language: str, source_id: str) -> str:
    digest = hashlib.sha256(
        f"{subqa}|{language}|{source_id}".encode("utf-8")
    ).hexdigest()[:12]
    return f"{subqa.replace('.', '')}-{language.upper()}-{digest}"


def build_qa11_cases(
    rows: Sequence[dict[str, str]],
    lex_index: dict[tuple[str, str], dict[str, str]],
) -> list[Case]:
    cases: list[Case] = []
    check_columns = [
        "target_triple_parity",
        "target_text_present",
        "lexicalisation_id_alignment",
        "lexicalisation_id_unique",
        "target_rdf_well_formed",
    ]
    for row in rows:
        failures = [
            column
            for column in check_columns
            if parse_bool(row.get(column)) is not True
        ]
        if not failures:
            risk = "Low"
            basis = "All five structural conditions pass."
        elif len(failures) >= 2 or any(
            column in failures
            for column in ("lexicalisation_id_alignment", "target_rdf_well_formed")
        ):
            risk = "High"
            basis = "Multiple or high-impact structural conditions fail: " + ", ".join(failures)
        else:
            risk = "Medium"
            basis = "One structural condition fails: " + failures[0]

        language = normalize_space(row.get("language"))
        record_key = normalize_space(row.get("record_key"))
        lexicalisation = lex_index.get((record_key, language), {})
        context = (
            f"Expected triples: {normalize_space(row.get('expected_triple_count'))}; "
            f"English IDs: {format_list(row.get('english_lids')) or '—'}; "
            f"Target IDs: {format_list(row.get('target_lids')) or '—'}; "
            f"Target lexicalisations in entry: {format_list(row.get('target_texts'), limit=3) or '—'}"
        )
        cases.append(Case(
            case_id=make_case_id("QA1.1", language, record_key),
            subqa="QA1.1",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=normalize_space(row.get("category")),
            focus=record_key,
            english_triples=format_list(row.get("source_triples")),
            target_triples=format_list(row.get("target_triples")),
            english_text=normalize_space(lexicalisation.get("source_text")),
            target_text=normalize_space(lexicalisation.get("variant_target_text")),
            comparison_context=context,
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals={
                column: parse_bool(row.get(column))
                for column in check_columns
            },
            uniqueness_token=f"QA11|{language}|{record_key}",
        ))
    return cases


def build_qa12_context_index(
    mappings: Sequence[dict[str, str]],
) -> dict[tuple[str, str, str], str]:
    grouped: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for row in mappings:
        key = (
            normalize_space(row.get("language")),
            normalize_space(row.get("component_type")),
            normalize_space(row.get("source_normalized")),
        )
        example = (
            f"{normalize_space(row.get('record_key'))}: "
            f"{normalize_space(row.get('source_raw'))} → "
            f"{normalize_space(row.get('target_raw'))}"
        )
        if example not in grouped[key] and len(grouped[key]) < 4:
            grouped[key].append(example)
    return {key: "\n".join(values) for key, values in grouped.items()}


def build_qa12_cases(
    profiles: Sequence[dict[str, str]],
    context_index: dict[tuple[str, str, str], str],
) -> list[Case]:
    cases: list[Case] = []
    for row in profiles:
        rate = parse_float(row.get("dominant_mapping_rate"))
        review = parse_bool(row.get("mapping_review_signal")) is True
        if not review:
            risk = "Low"
            basis = "One normalised target form is observed for the recurring source component."
        elif rate is not None and rate < 0.90:
            risk = "High"
            basis = "Multiple target forms are observed and the dominant form accounts for less than 90% of mappings."
        else:
            risk = "Medium"
            basis = "Multiple target forms are observed, but one form remains strongly dominant."

        language = normalize_space(row.get("language"))
        component_type = normalize_space(row.get("component_type"))
        source_normalized = normalize_space(row.get("source_normalized"))
        source_example = normalize_space(row.get("source_example")) or source_normalized
        target_forms = normalize_space(row.get("target_forms"))
        context = context_index.get((language, component_type, source_normalized), "")
        comparison = (
            f"Observed target forms and counts: {target_forms}\n"
            f"Example mappings:\n{context}"
        )
        cases.append(Case(
            case_id=make_case_id("QA1.2", language, f"{component_type}|{source_normalized}"),
            subqa="QA1.2",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=component_type,
            focus=source_example,
            english_triples="",
            target_triples="",
            english_text="",
            target_text="",
            comparison_context=comparison,
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals={
                "occurrences": parse_int(row.get("occurrences")),
                "target_form_count": parse_int(row.get("target_form_count")),
                "dominant_mapping_rate": rate,
                "mapping_review_signal": review,
            },
            uniqueness_token=f"QA12|{language}|{component_type}|{source_normalized}",
        ))
    return cases


def build_qa21_cases(
    rows: Sequence[dict[str, str]],
    thresholds: dict[str, dict[str, dict[float, float]]],
) -> list[Case]:
    metric_columns = ("text_similarity", "triple_similarity", "predicate_similarity")
    cases: list[Case] = []
    for row in rows:
        language = normalize_space(row.get("language"))
        below_01: list[str] = []
        below_05: list[str] = []
        below_10: list[str] = []
        signals: dict[str, Any] = {}
        for column in metric_columns:
            value = parse_float(row.get(column))
            signals[column] = value
            if value is None or language not in thresholds or column not in thresholds[language]:
                continue
            if value < thresholds[language][column][0.01]:
                below_01.append(column)
            if value < thresholds[language][column][0.05]:
                below_05.append(column)
            if value < thresholds[language][column][0.10]:
                below_10.append(column)

        literal = parse_float(row.get("source_target_literal_preservation"))
        literal_failure = literal is not None and literal < 1.0
        structured_flag = parse_bool(row.get("structured_predicate_review_flag")) is True
        if literal_failure or below_01 or len(below_05) >= 2:
            risk = "High"
            reasons = []
            if literal_failure:
                reasons.append("source literal preservation is below 1")
            if below_01:
                reasons.append("a semantic score is in the lowest 1%")
            if len(below_05) >= 2:
                reasons.append("at least two semantic scores are in the lowest 5%")
            basis = "; ".join(reasons) + "."
        elif structured_flag or below_10:
            risk = "Medium"
            reasons = []
            if structured_flag:
                reasons.append("the structured unchanged-predicate review flag is active")
            if below_10:
                reasons.append("at least one semantic score is in the lowest 10%")
            basis = "; ".join(reasons) + "."
        else:
            risk = "Low"
            basis = "No literal failure or lower-tail semantic review condition is present."

        signals.update({
            "source_target_literal_preservation": literal,
            "structured_predicate_review_flag": structured_flag,
            "below_1_percent": below_01,
            "below_5_percent": below_05,
            "below_10_percent": below_10,
        })
        focus = normalize_space(row.get("base_pair_key")) or normalize_space(row.get("variant_id"))
        cases.append(Case(
            case_id=make_case_id("QA2.1", language, focus),
            subqa="QA2.1",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=normalize_space(row.get("category")),
            focus=focus,
            english_triples=format_list(row.get("source_triples")),
            target_triples=format_list(row.get("variant_target_triples")),
            english_text=normalize_space(row.get("source_text")),
            target_text=normalize_space(row.get("variant_target_text")),
            comparison_context=(
                f"English lexicalisation ID: {normalize_space(row.get('source_lid'))}; "
                f"target lexicalisation ID: {normalize_space(row.get('variant_lid'))}"
            ),
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals=signals,
            uniqueness_token=f"LEX|{focus}",
        ))
    return cases


def build_qa22_cases(rows: Sequence[dict[str, str]]) -> list[Case]:
    cases: list[Case] = []
    for row in rows:
        coverage = parse_bool(row.get("english_calibrated_coverage_pass"))
        groundedness = parse_bool(row.get("english_calibrated_groundedness_pass"))
        literal = parse_float(row.get("target_text_literal_retention"))
        literal_failure = literal is not None and literal < 1.0
        failures = sum(value is False for value in (coverage, groundedness))
        undefined = sum(value is None for value in (coverage, groundedness))

        if literal_failure or failures >= 2:
            risk = "High"
            reasons = []
            if literal_failure:
                reasons.append("target literal retention is below 1")
            if failures >= 2:
                reasons.append("both English-calibrated support directions fail")
            basis = "; ".join(reasons) + "."
        elif failures == 1 or undefined > 0:
            risk = "Medium"
            basis = (
                "One English-calibrated support direction fails."
                if failures == 1
                else "At least one exact-size English calibration decision is unavailable."
            )
        else:
            risk = "Low"
            basis = "Coverage and groundedness pass; no target literal failure is observed."

        language = normalize_space(row.get("language"))
        focus = normalize_space(row.get("base_pair_key")) or normalize_space(row.get("variant_id"))
        cases.append(Case(
            case_id=make_case_id("QA2.2", language, focus),
            subqa="QA2.2",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=normalize_space(row.get("category")),
            focus=focus,
            english_triples=format_list(row.get("source_triples")),
            target_triples=format_list(row.get("variant_target_triples")),
            english_text=normalize_space(row.get("source_text")),
            target_text=normalize_space(row.get("variant_target_text")),
            comparison_context="Judge the target text only against the target triples; the English pair is supplied as alignment context.",
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals={
                "minimum_triple_coverage": parse_float(row.get("minimum_triple_coverage")),
                "minimum_text_groundedness": parse_float(row.get("minimum_text_groundedness")),
                "coverage_pass": coverage,
                "groundedness_pass": groundedness,
                "target_text_literal_retention": literal,
            },
            uniqueness_token=f"LEX|{focus}",
        ))
    return cases


def build_qa31_cases(rows: Sequence[dict[str, str]]) -> list[Case]:
    cases: list[Case] = []
    for row in rows:
        hard = parse_bool(row.get("hard_language_flag")) is True
        valid = parse_bool(row.get("valid_target_language")) is True
        if hard:
            risk = "High"
            basis = "At least one hard language diagnostic is active."
        elif not valid:
            risk = "Medium"
            basis = "The expected language does not reach the confidence-based valid-language decision, but no hard flag is active."
        else:
            risk = "Low"
            basis = "The expected target language is confidently predicted and no hard flag is active."

        language = normalize_space(row.get("language"))
        focus = normalize_space(row.get("base_pair_key")) or normalize_space(row.get("variant_id"))
        cases.append(Case(
            case_id=make_case_id("QA3.1", language, focus),
            subqa="QA3.1",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=normalize_space(row.get("category")),
            focus=focus,
            english_triples=format_list(row.get("source_triples")),
            target_triples=format_list(row.get("variant_target_triples")),
            english_text=normalize_space(row.get("source_text")),
            target_text=normalize_space(row.get("variant_target_text")),
            comparison_context="Assess language identity and substantial English leakage; do not score general fluency or style here.",
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals={
                "predicted_language": normalize_space(row.get("predicted_language")),
                "target_language_probability": parse_float(row.get("target_language_probability")),
                "valid_target_language": valid,
                "wrong_language_detected": parse_bool(row.get("wrong_language_detected")),
                "full_english_copy_detected": parse_bool(row.get("full_english_copy_detected")),
                "code_switch_detected": parse_bool(row.get("code_switch_detected")),
                "hard_language_flag": hard,
            },
            uniqueness_token=f"LEX|{focus}",
        ))
    return cases


def build_qa32_cases(
    rows: Sequence[dict[str, str]],
    thresholds: dict[str, dict[str, dict[float, float]]],
) -> list[Case]:
    cases: list[Case] = []
    for row in rows:
        language = normalize_space(row.get("language"))
        expansion = parse_float(row.get("expansion_ratio"))
        advantage = parse_float(row.get("aligned_reference_advantage"))
        language_thresholds = thresholds.get(language, {})
        expansion_thresholds = language_thresholds.get("expansion_ratio", {})
        advantage_thresholds = language_thresholds.get("aligned_reference_advantage", {})

        extreme_expansion = (
            expansion is not None
            and expansion_thresholds
            and (
                expansion < expansion_thresholds[0.01]
                or expansion > expansion_thresholds[0.99]
            )
        )
        unusual_expansion = (
            expansion is not None
            and expansion_thresholds
            and (
                expansion < expansion_thresholds[0.05]
                or expansion > expansion_thresholds[0.95]
            )
        )
        extreme_advantage = (
            advantage is not None
            and advantage_thresholds
            and advantage <= advantage_thresholds[0.01]
        )
        nonpositive_advantage = advantage is not None and advantage <= 0.0

        if extreme_advantage or (extreme_expansion and nonpositive_advantage):
            risk = "High"
            basis = "The aligned-reference advantage is in the extreme lower tail, or a non-positive advantage coincides with extreme length variation."
        elif nonpositive_advantage or unusual_expansion or advantage is None:
            risk = "Medium"
            basis = "The aligned-reference advantage is non-positive/unavailable, or the expansion ratio is outside the central 90% range."
        else:
            risk = "Low"
            basis = "The target favours its aligned English reference and its length lies within the central language-specific range."

        focus = normalize_space(row.get("base_pair_key")) or normalize_space(row.get("variant_id"))
        comparison_text = normalize_space(row.get("comparison_source_text"))
        context_parts = []
        if comparison_text:
            context_parts.append("Different English reference from the same RDF entry:\n" + comparison_text)
        else:
            context_parts.append("No different eligible English reference is available; judge whether the observed target realisation remains plausible.")
        cases.append(Case(
            case_id=make_case_id("QA3.2", language, focus),
            subqa="QA3.2",
            language=language,
            language_label=normalize_space(row.get("language_label")) or language,
            category=normalize_space(row.get("category")),
            focus=focus,
            english_triples=format_list(row.get("source_triples")),
            target_triples=format_list(row.get("variant_target_triples")),
            english_text=normalize_space(row.get("source_text")),
            target_text=normalize_space(row.get("variant_target_text")),
            comparison_context="\n".join(context_parts),
            automatic_risk=risk,
            automatic_basis=basis,
            automatic_signals={
                "expansion_ratio": expansion,
                "aligned_reference_eligible": parse_bool(row.get("aligned_reference_eligible")),
                "aligned_reference_similarity": parse_float(row.get("aligned_reference_similarity")),
                "comparison_reference_similarity": parse_float(row.get("comparison_reference_similarity")),
                "aligned_reference_advantage": advantage,
                "extreme_expansion": bool(extreme_expansion),
                "unusual_expansion": bool(unusual_expansion),
                "extreme_advantage": bool(extreme_advantage),
            },
            uniqueness_token=f"LEX|{focus}",
        ))
    return cases


def sample_cases(
    cases: Sequence[Case],
    *,
    per_subqa_language: int,
    seed: int,
) -> tuple[list[Case], list[dict[str, Any]]]:
    grouped: dict[tuple[str, str], list[Case]] = defaultdict(list)
    for case in cases:
        grouped[(case.subqa, case.language)].append(case)

    quota_template = {"Low": 4, "Medium": 3, "High": 3}
    if per_subqa_language != 10:
        # Preserve the approximate 40/30/30 allocation.
        low = max(1, round(per_subqa_language * 0.4))
        medium = max(0, round(per_subqa_language * 0.3))
        high = max(0, per_subqa_language - low - medium)
        quota_template = {"Low": low, "Medium": medium, "High": high}

    sampled: list[Case] = []
    summary: list[dict[str, Any]] = []
    used_lexicalisations: set[str] = set()

    for subqa in SUBQA_ORDER:
        languages = sorted({language for (candidate_subqa, language) in grouped if candidate_subqa == subqa})
        for language in languages:
            group = grouped[(subqa, language)]
            rng = random.Random(stable_seed(seed, subqa, language))
            by_risk: dict[str, list[Case]] = {risk: [] for risk in RISK_LEVELS}
            for case in group:
                by_risk[case.automatic_risk].append(case)
            for risk in RISK_LEVELS:
                rng.shuffle(by_risk[risk])

            chosen: list[Case] = []
            chosen_ids: set[str] = set()

            def choose_from(risk: str, count: int, *, avoid_reuse: bool) -> int:
                added = 0
                for case in by_risk[risk]:
                    if added >= count:
                        break
                    if case.case_id in chosen_ids:
                        continue
                    if avoid_reuse and case.uniqueness_token.startswith("LEX|") and case.uniqueness_token in used_lexicalisations:
                        continue
                    chosen.append(case)
                    chosen_ids.add(case.case_id)
                    added += 1
                return added

            # First respect the desired risk allocation and avoid repeated lexicalisations.
            for risk in RISK_LEVELS:
                choose_from(risk, quota_template[risk], avoid_reuse=True)

            # Fill shortages with unused cases, prioritising higher-risk strata.
            remaining = per_subqa_language - len(chosen)
            if remaining > 0:
                for risk in ("High", "Medium", "Low"):
                    if remaining <= 0:
                        break
                    added = choose_from(risk, remaining, avoid_reuse=True)
                    remaining -= added

            # If rare strata force reuse across sub-QAs, permit it only as a fallback.
            if remaining > 0:
                for risk in ("High", "Medium", "Low"):
                    if remaining <= 0:
                        break
                    added = choose_from(risk, remaining, avoid_reuse=False)
                    remaining -= added

            for case in chosen:
                if case.uniqueness_token.startswith("LEX|"):
                    used_lexicalisations.add(case.uniqueness_token)
            rng.shuffle(chosen)
            sampled.extend(chosen)

            sampled_counts = Counter(case.automatic_risk for case in chosen)
            available_counts = Counter(case.automatic_risk for case in group)
            summary.append({
                "subqa": subqa,
                "language": language,
                "available_low": available_counts.get("Low", 0),
                "available_medium": available_counts.get("Medium", 0),
                "available_high": available_counts.get("High", 0),
                "sampled_low": sampled_counts.get("Low", 0),
                "sampled_medium": sampled_counts.get("Medium", 0),
                "sampled_high": sampled_counts.get("High", 0),
                "sampled_total": len(chosen),
            })

    random.Random(seed).shuffle(sampled)
    return sampled, summary



# ---------------------------------------------------------------------------
# XLSX helpers (openpyxl)
# ---------------------------------------------------------------------------

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FONT = Font(bold=True, color="FFFFFF", size=14)
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True, color="17365D")
LIGHT_FILL = PatternFill("solid", fgColor="EEF5FB")
THIN_GREY = Side(style="thin", color="7F8C8D")
HAIR_GREY = Side(style="hair", color="D9D9D9")
HEADER_BORDER = Border(
    top=THIN_GREY,
    bottom=THIN_GREY,
    left=THIN_GREY,
    right=THIN_GREY,
)
BODY_BORDER = Border(
    top=HAIR_GREY,
    bottom=HAIR_GREY,
    left=HAIR_GREY,
    right=HAIR_GREY,
)


def new_workbook() -> Workbook:
    """Create an empty openpyxl workbook without the default sheet."""
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    return workbook


def write_matrix(sheet: Any, start_row: int, start_col: int, values: Sequence[Sequence[Any]]) -> None:
    """Write a rectangular list of rows starting at a 1-based cell position."""
    for row_offset, row_values in enumerate(values):
        for col_offset, value in enumerate(row_values):
            sheet.cell(
                row=start_row + row_offset,
                column=start_col + col_offset,
                value=value,
            )


def style_title(sheet: Any, cell_range: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.fill = TITLE_FILL
            cell.font = TITLE_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")


def style_header(sheet: Any, cell_range: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = HEADER_BORDER


def style_body(sheet: Any, cell_range: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BODY_BORDER


def add_instructions_sheet(wb: Workbook, *, blinded: bool) -> None:
    sheet = wb.create_sheet("Instructions")
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Simplified manual validation of the RDF-to-text audit"
    style_title(sheet, "A1:F1")
    sheet.row_dimensions[1].height = 28

    lines = [
        ["Purpose", "Provide a small human validation of the automatic audit without reproducing every metric-specific decision."],
        ["Unit", "Each row concerns one sub-QA only. Judge the high-level question shown in the Codebook."],
        ["Low", "No material issue is visible; the record or mapping is acceptable for this sub-QA."],
        ["Medium", "A possible or ambiguous issue warrants review, but may reflect legitimate variation or limited context."],
        ["High", "A clear substantive issue is visible for the focal sub-QA."],
        ["Cannot judge", "The supplied context is insufficient to make a reliable judgement."],
        ["Important", "Do not use the manual label to score unrelated properties. For example, QA3.1 concerns language identity/leakage, not general fluency."],
        [
            "Blinding",
            "Automatic risk levels and metric values are hidden from the annotator."
            if blinded
            else "This key workbook contains automatic risk levels and their metric-based rationale.",
        ],
    ]
    write_matrix(sheet, 3, 1, lines)
    for row in range(3, 3 + len(lines)):
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.cell(row=row, column=1).fill = LIGHT_FILL
        sheet.cell(row=row, column=1).alignment = Alignment(vertical="top")
        sheet.cell(row=row, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 90


def add_codebook_sheet(wb: Workbook) -> None:
    sheet = wb.create_sheet("Codebook")
    headers = ["Sub-QA", "Focus", "Manual question", "Low risk", "Medium risk", "High risk", "Generic example"]
    rows = []
    for subqa in SUBQA_ORDER:
        definition = SUBQA_DEFINITIONS[subqa]
        rows.append([
            subqa,
            definition["title"],
            definition["question"],
            definition["low"],
            definition["medium"],
            definition["high"],
            definition["example"],
        ])
    write_matrix(sheet, 1, 1, [headers] + rows)
    style_header(sheet, f"A1:G1")
    if rows:
        style_body(sheet, f"A2:G{len(rows) + 1}")
    for index, width in enumerate([10, 28, 48, 42, 42, 42, 48], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(2, len(rows) + 2):
        sheet.row_dimensions[row].height = 82
    sheet.freeze_panes = "A2"


def add_risk_rules_sheet(wb: Workbook, sampled: Sequence[Case]) -> None:
    del sampled  # Kept in the signature for API compatibility.
    sheet = wb.create_sheet("Automatic Risk Rules")
    headers = ["Sub-QA", "Low", "Medium", "High", "Role in the audit"]
    rows = [
        ["QA1.1", "All five structural checks pass.", "One lower-impact condition fails.", "Multiple failures or an ID-alignment/RDF-well-formedness failure.", "Deterministic risk mapping."],
        ["QA1.2", "One normalised target form.", "Multiple forms with dominant rate at least 0.90.", "Multiple forms with dominant rate below 0.90.", "Review mapping; variation is not automatically an error."],
        ["QA2.1", "No literal failure or lower-tail semantic signal.", "Structured predicate flag or at least one score in the lowest 10%.", "Literal failure, a score in the lowest 1%, or at least two scores in the lowest 5%.", "Language-relative prioritisation, not a universal correctness threshold."],
        ["QA2.2", "Coverage and groundedness pass; no target literal failure.", "One pass fails or calibration is unavailable.", "Both passes fail or target literal retention is below 1.", "Combines omission/addition review directions."],
        ["QA3.1", "Confident valid target language and no hard flag.", "Uncertain target-language decision without a hard flag.", "Wrong-language, exact-copy, or code-switch hard flag.", "High risk corresponds to explicit language-related diagnostics."],
        ["QA3.2", "Positive aligned-reference advantage and central length range.", "Non-positive/unavailable advantage or length outside the central 90%.", "Extreme lower-tail advantage, or non-positive advantage with extreme length.", "Describes risk that variation is not legitimate; it is not a direct error metric."],
    ]
    write_matrix(sheet, 1, 1, [headers] + rows)
    style_header(sheet, "A1:E1")
    style_body(sheet, f"A2:E{len(rows) + 1}")
    for column, width in zip("ABCDE", [10, 45, 50, 55, 45]):
        sheet.column_dimensions[column].width = width
    for row in range(2, len(rows) + 2):
        sheet.row_dimensions[row].height = 75
    sheet.freeze_panes = "A2"


def add_sampling_summary_sheet(wb: Workbook, summary: Sequence[dict[str, Any]]) -> None:
    sheet = wb.create_sheet("Sampling Summary")
    headers = [
        "Sub-QA", "Language", "Available low", "Available medium", "Available high",
        "Sampled low", "Sampled medium", "Sampled high", "Sampled total",
    ]
    rows = [[
        row["subqa"], row["language"], row["available_low"], row["available_medium"], row["available_high"],
        row["sampled_low"], row["sampled_medium"], row["sampled_high"], row["sampled_total"],
    ] for row in summary]
    write_matrix(sheet, 1, 1, [headers] + rows)
    style_header(sheet, "A1:I1")
    if rows:
        style_body(sheet, f"A2:I{len(rows) + 1}")
    for column in "ABCDEFGHI":
        sheet.column_dimensions[column].width = 16
    sheet.freeze_panes = "A2"


def add_case_sheet(wb: Workbook, subqa: str, cases: Sequence[Case], *, blinded: bool) -> None:
    sheet = wb.create_sheet(SHEET_NAMES[subqa])
    definition = SUBQA_DEFINITIONS[subqa]
    final_column = "K" if blinded else "L"

    sheet.merge_cells(f"A1:{final_column}1")
    sheet["A1"] = f"{subqa}: {definition['title']}"
    style_title(sheet, f"A1:{final_column}1")

    sheet.merge_cells(f"A2:{final_column}2")
    sheet["A2"] = definition["question"]
    for row in sheet[f"A2:{final_column}2"]:
        for cell in row:
            cell.fill = LIGHT_FILL
            cell.font = Font(italic=True, color="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    if blinded:
        headers = ANNOTATOR_HEADERS
        rows = [case.annotator_row() for case in cases]
    else:
        headers = ANNOTATOR_HEADERS[:-2] + KEY_EXTRA_HEADERS
        rows = [case.key_row() for case in cases]

    start_row = 4
    end_row = start_row + len(rows)
    write_matrix(sheet, start_row, 1, [headers] + rows)
    style_header(sheet, f"A{start_row}:{final_column}{start_row}")
    if rows:
        style_body(sheet, f"A{start_row + 1}:{final_column}{end_row}")
        for row in range(start_row + 1, end_row + 1):
            sheet.row_dimensions[row].height = 105

    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A{start_row}:{final_column}{end_row}"

    widths = {
        "A": 24, "B": 12, "C": 18, "D": 36, "E": 46, "F": 46,
        "G": 48, "H": 48, "I": 52,
    }
    if blinded:
        widths.update({"J": 17, "K": 45})
    else:
        widths.update({"J": 17, "K": 60, "L": 65})
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    if blinded and rows:
        risk_cells = f"J{start_row + 1}:J{end_row}"
        validation = DataValidation(
            type="list",
            formula1='"Low,Medium,High,Cannot judge"',
            allow_blank=True,
        )
        validation.error = "Choose Low, Medium, High, or Cannot judge."
        validation.errorTitle = "Invalid manual risk"
        validation.prompt = "Select one manual risk level."
        validation.promptTitle = "Manual risk"
        sheet.add_data_validation(validation)
        validation.add(risk_cells)

        conditional_styles = {
            "Low": ("E2F0D9", "375623"),
            "Medium": ("FFF2CC", "7F6000"),
            "High": ("F4CCCC", "990000"),
            "Cannot judge": ("D9D9D9", "404040"),
        }
        first_data_row = start_row + 1
        for label, (fill_color, font_color) in conditional_styles.items():
            sheet.conditional_formatting.add(
                risk_cells,
                FormulaRule(
                    formula=[f'$J{first_data_row}="{label}"'],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color),
                ),
            )


def build_workbooks(
    sampled: Sequence[Case],
    summary: Sequence[dict[str, Any]],
    output_dir: Path,
) -> tuple[Path, Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[Case]] = defaultdict(list)
    for case in sampled:
        grouped[case.subqa].append(case)
    for subqa in grouped:
        grouped[subqa].sort(key=lambda case: case.case_id)

    annotator = new_workbook()
    add_instructions_sheet(annotator, blinded=True)
    add_codebook_sheet(annotator)
    for subqa in SUBQA_ORDER:
        add_case_sheet(annotator, subqa, grouped.get(subqa, []), blinded=True)

    key = new_workbook()
    add_instructions_sheet(key, blinded=False)
    add_codebook_sheet(key)
    add_risk_rules_sheet(key, sampled)
    add_sampling_summary_sheet(key, summary)
    for subqa in SUBQA_ORDER:
        add_case_sheet(key, subqa, grouped.get(subqa, []), blinded=False)

    annotator_path = output_dir / "manual_audit_annotator.xlsx"
    key_path = output_dir / "manual_audit_key.xlsx"
    key_csv_path = output_dir / "manual_audit_key.csv"
    summary_csv_path = output_dir / "manual_audit_sampling_summary.csv"

    annotator.save(annotator_path)
    key.save(key_path)

    key_rows = [case.combined_key_dict() for case in sampled]
    key_fields = list(key_rows[0].keys()) if key_rows else [
        "case_id", "subqa", "language", "language_label", "category",
        "record_or_component", "english_triples", "target_triples",
        "aligned_english_text", "target_text", "comparison_or_context",
        "automatic_risk", "automatic_basis", "automatic_signals",
    ]
    write_csv_rows(key_csv_path, key_rows, key_fields)
    write_csv_rows(
        summary_csv_path,
        summary,
        [
            "subqa", "language", "available_low", "available_medium", "available_high",
            "sampled_low", "sampled_medium", "sampled_high", "sampled_total",
        ],
    )
    return annotator_path, key_path, key_csv_path, summary_csv_path


def inspect_annotation_rows(path: Path, key_rows: Sequence[dict[str, str]]) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    counts = Counter(row["subqa"] for row in key_rows)
    annotations: list[dict[str, str]] = []

    for subqa in SUBQA_ORDER:
        count = counts.get(subqa, 0)
        if count == 0:
            continue
        sheet_name = SHEET_NAMES[subqa]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing annotation sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        start_row = 4
        headers = [normalize_space(sheet.cell(start_row, column).value) for column in range(1, 12)]
        header_to_col = {header: index + 1 for index, header in enumerate(headers)}
        required = {"Case ID", "Manual risk", "Manual notes"}
        missing = required - set(header_to_col)
        if missing:
            raise ValueError(
                f"Sheet {sheet_name} is missing required columns: {sorted(missing)}"
            )

        for row_index in range(start_row + 1, start_row + count + 1):
            case_id = normalize_space(sheet.cell(row_index, header_to_col["Case ID"]).value)
            if not case_id:
                continue
            annotations.append({
                "case_id": case_id,
                "manual_risk": normalize_space(
                    sheet.cell(row_index, header_to_col["Manual risk"]).value
                ),
                "manual_notes": normalize_space(
                    sheet.cell(row_index, header_to_col["Manual notes"]).value
                ),
            })

    workbook.close()
    return annotations


def safe_rate(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def evaluation_summary(rows: Sequence[dict[str, Any]], group_fields: Sequence[str]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = tuple(normalize_space(row.get(field)) for field in group_fields)
        grouped[key].append(row)

    summaries: list[dict[str, Any]] = []
    for key, group in sorted(grouped.items()):
        usable = [row for row in group if row["manual_risk"] in RISK_TO_NUMBER]
        cannot_judge = sum(row["manual_risk"] == "Cannot judge" for row in group)
        exact = sum(row["manual_risk"] == row["automatic_risk"] for row in usable)
        distances = [
            abs(RISK_TO_NUMBER[row["manual_risk"]] - RISK_TO_NUMBER[row["automatic_risk"]])
            for row in usable
        ]
        adjacent = sum(distance <= 1 for distance in distances)

        tp = fp = fn = tn = 0
        for row in usable:
            auto_priority = RISK_TO_NUMBER[row["automatic_risk"]] >= 1
            manual_priority = RISK_TO_NUMBER[row["manual_risk"]] >= 1
            if auto_priority and manual_priority:
                tp += 1
            elif auto_priority and not manual_priority:
                fp += 1
            elif not auto_priority and manual_priority:
                fn += 1
            else:
                tn += 1
        precision = safe_rate(tp, tp + fp)
        recall = safe_rate(tp, tp + fn)
        f1 = (
            2 * precision * recall / (precision + recall)
            if precision is not None and recall is not None and precision + recall > 0
            else None
        )
        auto_low_confirmation = safe_rate(tn, tn + fn)

        summary = {
            **{field: value for field, value in zip(group_fields, key)},
            "total_cases": len(group),
            "usable_cases": len(usable),
            "cannot_judge": cannot_judge,
            "exact_agreement": safe_rate(exact, len(usable)),
            "within_one_level": safe_rate(adjacent, len(usable)),
            "mean_absolute_level_difference": (
                sum(distances) / len(distances) if distances else None
            ),
            "priority_precision": precision,
            "priority_recall": recall,
            "priority_f1": f1,
            "automatic_low_confirmation": auto_low_confirmation,
            "tp": tp,
            "fp": fp,
            "fn": fn,
            "tn": tn,
        }
        summaries.append(summary)
    return summaries


def write_evaluation_workbook(
    rows: Sequence[dict[str, Any]],
    output_dir: Path,
) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    overall = evaluation_summary(rows, [])
    by_subqa = evaluation_summary(rows, ["subqa"])
    by_language = evaluation_summary(rows, ["language_label"])
    by_subqa_language = evaluation_summary(rows, ["subqa", "language_label"])

    summary_fields = [
        "total_cases", "usable_cases", "cannot_judge", "exact_agreement",
        "within_one_level", "mean_absolute_level_difference", "priority_precision",
        "priority_recall", "priority_f1", "automatic_low_confirmation", "tp", "fp", "fn", "tn",
    ]

    workbook = new_workbook()
    instructions = workbook.create_sheet("Interpretation")
    instructions.merge_cells("A1:F1")
    instructions["A1"] = "Simplified manual-audit evaluation"
    style_title(instructions, "A1:F1")
    interpretation_rows = [
        ["Exact agreement", "Automatic and manual risk levels are identical."],
        ["Within one level", "Automatic and manual labels differ by no more than one risk level."],
        ["Priority precision", "Among automatic Medium/High cases, the proportion also judged Medium/High manually."],
        ["Priority recall", "Among manually judged Medium/High cases, the proportion prioritised automatically."],
        ["Automatic-low confirmation", "Among automatic Low cases, the proportion also judged Low manually."],
        ["Cannot judge", "Excluded from agreement and priority calculations."],
    ]
    write_matrix(instructions, 3, 1, interpretation_rows)
    for row in range(3, 3 + len(interpretation_rows)):
        instructions.cell(row=row, column=1).font = Font(bold=True)
        instructions.cell(row=row, column=1).fill = LIGHT_FILL
        instructions.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 90

    def add_summary(name: str, data: Sequence[dict[str, Any]], grouping: Sequence[str]) -> None:
        sheet = workbook.create_sheet(name)
        headers = [*grouping, *summary_fields]
        values = [[row.get(column) for column in headers] for row in data]
        write_matrix(sheet, 1, 1, [headers] + values)
        end_col = get_column_letter(len(headers))
        style_header(sheet, f"A1:{end_col}1")
        if values:
            style_body(sheet, f"A2:{end_col}{len(values) + 1}")
        for index, header in enumerate(headers, start=1):
            column = get_column_letter(index)
            sheet.column_dimensions[column].width = 18 if header not in grouping else 16
            if header in {
                "exact_agreement", "within_one_level", "priority_precision", "priority_recall",
                "priority_f1", "automatic_low_confirmation", "mean_absolute_level_difference",
            }:
                for row in range(2, len(values) + 2):
                    sheet.cell(row=row, column=index).number_format = "0.000"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{end_col}{len(values) + 1}"

    add_summary("Overall", overall, [])
    add_summary("By Sub-QA", by_subqa, ["subqa"])
    add_summary("By Language", by_language, ["language_label"])
    add_summary("By Sub-QA Language", by_subqa_language, ["subqa", "language_label"])

    confusion = Counter(
        (row["automatic_risk"], row["manual_risk"])
        for row in rows
        if row["manual_risk"] in RISK_TO_NUMBER
    )
    confusion_sheet = workbook.create_sheet("Confusion Matrix")
    confusion_values = [["Automatic \\ Manual", *RISK_LEVELS]]
    for auto in RISK_LEVELS:
        confusion_values.append([
            auto,
            *[confusion.get((auto, manual), 0) for manual in RISK_LEVELS],
        ])
    write_matrix(confusion_sheet, 1, 1, confusion_values)
    style_header(confusion_sheet, "A1:D1")
    style_header(confusion_sheet, "A2:A4")
    for column in "ABCD":
        confusion_sheet.column_dimensions[column].width = 20

    disagreements = [
        row for row in rows
        if row["manual_risk"] in RISK_TO_NUMBER
        and row["manual_risk"] != row["automatic_risk"]
    ]
    disagreement_sheet = workbook.create_sheet("Disagreements")
    disagreement_headers = [
        "case_id", "subqa", "language_label", "automatic_risk", "manual_risk",
        "automatic_basis", "manual_notes", "record_or_component",
    ]
    disagreement_values = [
        [row.get(column, "") for column in disagreement_headers]
        for row in disagreements
    ]
    write_matrix(disagreement_sheet, 1, 1, [disagreement_headers] + disagreement_values)
    style_header(disagreement_sheet, "A1:H1")
    if disagreement_values:
        style_body(disagreement_sheet, f"A2:H{len(disagreement_values) + 1}")
    for column, width in zip("ABCDEFGH", [24, 10, 14, 16, 16, 60, 50, 40]):
        disagreement_sheet.column_dimensions[column].width = width
    disagreement_sheet.freeze_panes = "A2"

    output_path = output_dir / "manual_audit_evaluation.xlsx"
    workbook.save(output_path)

    summary_csv = output_dir / "manual_audit_evaluation_by_subqa_language.csv"
    write_csv_rows(summary_csv, by_subqa_language, ["subqa", "language_label", *summary_fields])
    disagreements_csv = output_dir / "manual_audit_disagreements.csv"
    write_csv_rows(disagreements_csv, disagreements, disagreement_headers)
    return output_path, summary_csv, disagreements_csv


def command_generate(args: argparse.Namespace) -> None:
    results_root = Path(args.results_root).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    paths = discover_paths(results_root)

    qa11_rows = read_csv_rows(paths.qa11)
    qa12_profiles = read_csv_rows(paths.qa12_profiles)
    qa12_mappings = read_csv_rows(paths.qa12_mappings)
    qa2_rows = read_csv_rows(paths.qa2)
    qa3_rows = read_csv_rows(paths.qa3)

    lex_index = first_lexicalisation_index(qa2_rows)
    qa21_thresholds = thresholds_by_language(
        qa2_rows,
        ("text_similarity", "triple_similarity", "predicate_similarity"),
        (0.01, 0.05, 0.10),
    )
    qa32_thresholds = thresholds_by_language(
        qa3_rows,
        ("expansion_ratio", "aligned_reference_advantage"),
        (0.01, 0.05, 0.95, 0.99),
    )

    cases = []
    cases.extend(build_qa11_cases(qa11_rows, lex_index))
    cases.extend(build_qa12_cases(qa12_profiles, build_qa12_context_index(qa12_mappings)))
    cases.extend(build_qa21_cases(qa2_rows, qa21_thresholds))
    cases.extend(build_qa22_cases(qa2_rows))
    cases.extend(build_qa31_cases(qa3_rows))
    cases.extend(build_qa32_cases(qa3_rows, qa32_thresholds))

    sampled, summary = sample_cases(
        cases,
        per_subqa_language=args.per_subqa_language,
        seed=args.seed,
    )
    outputs = build_workbooks(sampled, summary, output_dir)

    threshold_path = output_dir / "manual_audit_thresholds.json"
    threshold_path.write_text(
        json.dumps(
            {
                "qa21_language_percentiles": qa21_thresholds,
                "qa32_language_percentiles": qa32_thresholds,
                "seed": args.seed,
                "per_subqa_language": args.per_subqa_language,
                "total_cases": len(sampled),
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print(f"Generated {len(sampled)} manual-audit cases.")
    for path in outputs:
        print(path)
    print(threshold_path)


def command_evaluate(args: argparse.Namespace) -> None:
    annotated_path = Path(args.annotated).expanduser().resolve()
    key_csv_path = Path(args.key_csv).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    key_rows = read_csv_rows(key_csv_path)
    annotations = inspect_annotation_rows(annotated_path, key_rows)
    annotation_index = {row["case_id"]: row for row in annotations}

    merged: list[dict[str, Any]] = []
    invalid_labels: list[tuple[str, str]] = []
    missing_annotations: list[str] = []
    for key in key_rows:
        case_id = key["case_id"]
        annotation = annotation_index.get(case_id)
        if annotation is None:
            missing_annotations.append(case_id)
            manual_risk = ""
            manual_notes = ""
        else:
            manual_risk = annotation["manual_risk"]
            manual_notes = annotation["manual_notes"]
        if manual_risk and manual_risk not in MANUAL_LEVELS:
            invalid_labels.append((case_id, manual_risk))
        merged.append({
            **key,
            "manual_risk": manual_risk,
            "manual_notes": manual_notes,
        })

    if missing_annotations:
        raise ValueError(
            f"{len(missing_annotations)} case IDs from the key are missing from the annotated workbook. "
            f"First missing IDs: {missing_annotations[:5]}"
        )
    if invalid_labels:
        raise ValueError(
            "Invalid manual risk labels. Use Low, Medium, High, or Cannot judge. "
            f"Examples: {invalid_labels[:5]}"
        )
    blank = [row["case_id"] for row in merged if not row["manual_risk"]]
    if blank and not args.allow_incomplete:
        raise ValueError(
            f"{len(blank)} cases have no manual risk label. Complete them or use --allow-incomplete. "
            f"First blank IDs: {blank[:5]}"
        )

    usable_merged = [row for row in merged if row["manual_risk"]]
    outputs = write_evaluation_workbook(usable_merged, output_dir)
    print(f"Evaluated {len(usable_merged)} labelled cases.")
    for path in outputs:
        print(path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    generate = subparsers.add_parser("generate", help="Generate blinded annotator and automatic-key files.")
    generate.add_argument("--results-root", required=True, help="Root containing QA1_stress_aligned, QA2_stress_aligned, and QA3_stress_aligned.")
    generate.add_argument("--output-dir", required=True, help="Directory for the generated manual-audit files.")
    generate.add_argument("--per-subqa-language", type=int, default=10, help="Target number of cases per sub-QA and language (default: 10; total target: 120).")
    generate.add_argument("--seed", type=int, default=42)
    generate.set_defaults(func=command_generate)

    evaluate = subparsers.add_parser("evaluate", help="Evaluate a completed annotator workbook against the automatic risk key.")
    evaluate.add_argument("--annotated", required=True, help="Completed copy of manual_audit_annotator.xlsx.")
    evaluate.add_argument("--key-csv", required=True, help="manual_audit_key.csv generated with the workbook.")
    evaluate.add_argument("--output-dir", required=True, help="Directory for evaluation outputs.")
    evaluate.add_argument("--allow-incomplete", action="store_true", help="Evaluate only completed rows when some labels are blank.")
    evaluate.set_defaults(func=command_evaluate)

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()